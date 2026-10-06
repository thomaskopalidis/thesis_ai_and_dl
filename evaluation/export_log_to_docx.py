"""
Converts THESIS_LOG.md into a .docx file for pasting into the thesis.

Usage:
  pip install python-docx markdown-it-py
  python evaluation/export_log_to_docx.py
  # produces THESIS_LOG.docx next to the .md file

Keeps formatting simple on purpose: headings, bullet lists, numbered lists,
tables, inline code, and fenced code blocks. Images and complex HTML are
ignored.
"""
from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
IN_PATH = PROJECT_ROOT / "THESIS_LOG.md"
OUT_PATH = PROJECT_ROOT / "THESIS_LOG.docx"


def main():
    try:
        from docx import Document
        from docx.shared import Pt
        from markdown_it import MarkdownIt
    except ImportError:
        print(
            "[ERROR] Missing dependencies. Install with:\n"
            "    pip install python-docx markdown-it-py",
            file=sys.stderr,
        )
        sys.exit(1)

    if not IN_PATH.is_file():
        print(f"[ERROR] {IN_PATH} not found.", file=sys.stderr)
        sys.exit(1)

    md = MarkdownIt("commonmark", {"html": False}).enable("table")
    tokens = md.parse(IN_PATH.read_text(encoding="utf-8"))

    doc = Document()
    # Make default font a bit more readable
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(tokens):
        t = tokens[i]

        if t.type == "heading_open":
            level = int(t.tag[1])  # h1..h6
            text = tokens[i + 1].content
            doc.add_heading(text, level=min(level, 4))
            i += 3  # heading_open, inline, heading_close
            continue

        if t.type == "paragraph_open":
            text = tokens[i + 1].content
            p = doc.add_paragraph()
            _add_inline(p, text)
            i += 3
            continue

        if t.type == "bullet_list_open":
            i = _consume_list(doc, tokens, i, numbered=False)
            continue

        if t.type == "ordered_list_open":
            i = _consume_list(doc, tokens, i, numbered=True)
            continue

        if t.type == "fence" or t.type == "code_block":
            p = doc.add_paragraph()
            run = p.add_run(t.content.rstrip())
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue

        if t.type == "table_open":
            i = _consume_table(doc, tokens, i)
            continue

        if t.type == "blockquote_open":
            # Render as italic paragraph
            end = i + 1
            while end < len(tokens) and tokens[end].type != "blockquote_close":
                end += 1
            text_parts = []
            for j in range(i + 1, end):
                if tokens[j].type == "inline":
                    text_parts.append(tokens[j].content)
            p = doc.add_paragraph()
            r = p.add_run("  " + " ".join(text_parts))
            r.italic = True
            i = end + 1
            continue

        if t.type == "hr":
            doc.add_paragraph("─" * 40)
            i += 1
            continue

        i += 1

    doc.save(OUT_PATH)
    print(f"[OK] Wrote {OUT_PATH}")


def _add_inline(paragraph, text: str):
    """Very small inline-markup renderer: `code`, **bold**, *italic*."""
    import re
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        tok = m.group(0)
        if tok.startswith("`"):
            r = paragraph.add_run(tok.strip("`"))
            r.font.name = "Consolas"
        elif tok.startswith("**"):
            r = paragraph.add_run(tok.strip("*"))
            r.bold = True
        else:
            r = paragraph.add_run(tok.strip("*"))
            r.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _consume_list(doc, tokens, i, numbered: bool) -> int:
    close = "ordered_list_close" if numbered else "bullet_list_close"
    style = "List Number" if numbered else "List Bullet"
    i += 1
    while i < len(tokens) and tokens[i].type != close:
        if tokens[i].type == "list_item_open":
            # find the inline content inside this item
            j = i + 1
            text = ""
            while j < len(tokens) and tokens[j].type != "list_item_close":
                if tokens[j].type == "inline":
                    text = tokens[j].content
                    break
                j += 1
            p = doc.add_paragraph(style=style)
            _add_inline(p, text)
            # advance past list_item_close
            while i < len(tokens) and tokens[i].type != "list_item_close":
                i += 1
        i += 1
    return i + 1  # skip the list_close


def _consume_table(doc, tokens, i) -> int:
    # Collect rows of inline cells.
    rows: list[list[str]] = []
    current: list[str] = []
    i += 1
    while i < len(tokens) and tokens[i].type != "table_close":
        if tokens[i].type == "tr_open":
            current = []
        elif tokens[i].type in ("th_open", "td_open"):
            # Next token is inline
            if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                current.append(tokens[i + 1].content)
            else:
                current.append("")
        elif tokens[i].type == "tr_close":
            rows.append(current)
        i += 1

    if not rows:
        return i + 1

    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline(p, cell_text)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    return i + 1


if __name__ == "__main__":
    main()
